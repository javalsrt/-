from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import re


def set_font(run, name='宋体', size=12, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)


def add_heading(doc, text, level=1):
    para = doc.add_heading(level=level)
    run = para.add_run(text)
    sizes = {1: 18, 2: 16, 3: 14}
    set_font(run, '黑体', sizes.get(level, 12), True)
    if level == 1:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return para


def add_para(doc, text, bold=False):
    para = doc.add_paragraph()
    para.paragraph_format.first_line_indent = Inches(0.44)
    para.paragraph_format.line_spacing = 1.5
    run = para.add_run(text)
    set_font(run, '宋体', 12, bold)


def add_code(doc, lines):
    for line in lines:
        para = doc.add_paragraph()
        para.paragraph_format.left_indent = Inches(0.3)
        run = para.add_run(line)
        set_font(run, 'Courier New', 10)


def md_to_docx(md_path, docx_path):
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = '宋体'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    style.font.size = Pt(12)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code = False
    code_lines = []
    table_lines = []

    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\n')

        if line.strip().startswith('```'):
            if in_code:
                add_code(doc, code_lines)
                code_lines = []
            in_code = not in_code
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if line.strip().startswith('|'):
            table_lines.append(line)
            i += 1
            continue
        elif table_lines:
            rows = []
            for tl in table_lines:
                cells = [c.strip() for c in tl.strip().split('|')]
                cells = [c for c in cells if c]
                if cells and not all(set(c) <= set(' -') for c in cells):
                    rows.append(cells)
            if rows:
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = 'Table Grid'
                for r, row_cells in enumerate(rows):
                    for c, text in enumerate(row_cells):
                        cell = table.rows[r].cells[c]
                        cell.text = text
                        for p in cell.paragraphs:
                            for run in p.runs:
                                set_font(run, '宋体', 10)
            table_lines = []
            continue

        m = re.match(r'^(#{1,4})\s+(.*)', line)
        if m:
            add_heading(doc, m.group(2).strip(), len(m.group(1)))
            i += 1
            continue

        if line.strip() == '---':
            i += 1
            continue

        if line.strip():
            add_para(doc, line.strip())
        i += 1

    doc.save(docx_path)
    print('saved', docx_path)


md_to_docx(
    r'C:\Users\jay\Desktop\作业\小林web大作业\综合实训报告.md',
    r'C:\Users\jay\Desktop\作业\小林web大作业\综合实训报告.docx'
)
